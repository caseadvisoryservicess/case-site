#!/usr/bin/env python3
"""Проверка выгруженных .xlsx и .pptx настоящими библиотеками.

Файлы CASE OS собирает сам, без сторонних библиотек в браузере, поэтому мало
убедиться, что файл скачался: он должен ОТКРЫВАТЬСЯ. openpyxl и python-pptx
разбирают тот же формат, что Excel и PowerPoint, — если структура битая,
здесь это видно сразу.

    pip install openpyxl python-pptx
    python3 check_office_files.py docs/qa/tools/exporttest
"""
import sys, os, glob, zipfile
from xml.etree import ElementTree

failed = 0


def ck(name, cond, detail=None):
    global failed
    print(('OK  ' if cond else '!!  ') + name + ('' if detail is None else ' — ' + str(detail)))
    if not cond:
        failed += 1


def check_zip_xml(path):
    """Каждая часть пакета должна быть валидным XML — иначе Office молча откажет."""
    bad = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith('.xml') or n.endswith('.rels'):
                try:
                    ElementTree.fromstring(z.read(n))
                except Exception as e:
                    bad.append(n + ': ' + str(e)[:60])
    return bad


def check_xlsx(path):
    import openpyxl
    ck('Excel: файл открывается openpyxl', True) if False else None
    bad = check_zip_xml(path)
    ck('Excel: все части пакета — валидный XML', not bad, '; '.join(bad) or 'ок')
    wb = openpyxl.load_workbook(path)
    ck('Excel: файл открывается', True, os.path.basename(path))
    names = wb.sheetnames
    want = ['Сводка', 'Население', 'Население по районам', 'Районы города', 'Радиусы',
            'Бизнес-центры', 'Медицина', 'Городские объекты', 'Метро', 'Рынок города']
    ck('Excel: все десять листов на месте', names == want, names)

    sv = wb['Сводка']
    vals = [str(c.value) for row in sv.iter_rows() for c in row if c.value is not None]
    ck('Excel: в сводке есть название проекта', any('Samsung' in v for v in vals), vals[:6])
    ck('Excel: в сводке есть координаты', any('41.3' in v for v in vals))
    ck('Excel: в сводке есть скоринг', any('Потенциал' in v for v in vals),
       [v for v in vals if 'Потенциал' in v][:4])
    ck('Excel: формула скоринга описана', any('0,55' in v for v in vals))

    rd = wb['Радиусы']
    head = [c.value for c in next(rd.iter_rows(min_row=1, max_row=1))]
    ck('Excel: шапка радиусов полная',
       head[:4] == ['Радиус, м', 'Население', 'Бизнес-центры', 'Медицина'] and 'F&B' in head and 'Образование' in head,
       head)
    body = list(rd.iter_rows(min_row=2, values_only=True))
    ck('Excel: строк по радиусам не меньше пяти', len(body) >= 5, len(body))
    nums = [r[0] for r in body]
    ck('Excel: радиусы — числа, а не текст', all(isinstance(x, (int, float)) for x in nums), nums)
    fnb_col = head.index('F&B')
    ck('Excel: F&B посчитан', any(isinstance(r[fnb_col], (int, float)) for r in body),
       [r[fnb_col] for r in body])

    # --- население: плотность и прирост по кольцам
    pop = wb['Население']
    ph = [c.value for c in next(pop.iter_rows(min_row=1, max_row=1))]
    ck('Excel: в населении есть плотность и прирост',
       'Плотность, чел/км²' in ph and 'Прирост в кольце' in ph and 'Доля населения района, %' in ph, ph)
    pbody = [r for r in pop.iter_rows(min_row=2, values_only=True) if isinstance(r[0], (int, float))]
    ck('Excel: строки населения посчитаны', len(pbody) >= 5, len(pbody))
    dens = [r[ph.index('Плотность, чел/км²')] for r in pbody]
    ck('Excel: плотность — числа', all(isinstance(x, (int, float)) for x in dens), dens)
    area = [r[ph.index('Площадь круга, км²')] for r in pbody]
    ck('Excel: площадь круга растёт с радиусом', area == sorted(area) and area[0] > 0, area)

    # --- разрез по районам
    bd = wb['Население по районам']
    bh = [c.value for c in next(bd.iter_rows(min_row=1, max_row=1))]
    ck('Excel: разрез населения по районам', bh == ['Район', 'Жителей в 1 км', 'Жителей в 3 км'], bh)

    # --- все районы города
    ad = wb['Районы города']
    arows = list(ad.iter_rows(min_row=2, values_only=True))
    ck('Excel: справочник районов города заполнен', len(arows) >= 10, f'районов: {len(arows)}')
    ck('Excel: у районов есть население и плотность',
       all(isinstance(r[1], (int, float)) and isinstance(r[3], (int, float)) for r in arows[:5]),
       arows[:2])
    ck('Excel: район проекта отмечен', any((r[4] or '') == 'да' for r in arows),
       [r[0] for r in arows if (r[4] or '') == 'да'])

    # --- бизнес-центры
    bc = wb['Бизнес-центры']
    bch = [c.value for c in next(bc.iter_rows(min_row=1, max_row=1))]
    ck('Excel: у конкурентов есть цена, класс и адрес',
       'Ставка, $/м²/мес' in bch and 'Класс' in bch and 'Адрес' in bch, bch)
    bcrows = [r for r in bc.iter_rows(min_row=2, values_only=True) if r[0]]
    ck('Excel: список конкурентов не пуст', len(bcrows) >= 1, f'строк: {len(bcrows)}')

    # --- метро
    mt = wb['Метро']
    mh = [c.value for c in next(mt.iter_rows(min_row=1, max_row=1))]
    ck('Excel: метро с расстоянием и временем пешком',
       'Расстояние, м' in mh and any('пешком' in str(x).lower() for x in mh), mh)


def check_pptx(path):
    from pptx import Presentation
    from pptx.util import Emu
    bad = check_zip_xml(path)
    ck('PowerPoint: все части пакета — валидный XML', not bad, '; '.join(bad) or 'ок')
    pr = Presentation(path)
    ck('PowerPoint: файл открывается', True, os.path.basename(path))
    ck('PowerPoint: формат 16:9', abs(pr.slide_width / pr.slide_height - 16 / 9) < 0.01,
       f'{pr.slide_width}×{pr.slide_height}')
    n = len(pr.slides.__iter__.__self__._sldIdLst)
    ck('PowerPoint: слайдов не меньше восьми', n >= 8, n)

    texts, tables = [], 0
    for sl in pr.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, 'has_table', False) and sh.has_table:
                tables += 1
    all_text = '\n'.join(texts)
    ck('PowerPoint: титул с названием проекта', 'Samsung' in all_text, texts[:3])
    ck('PowerPoint: есть слайд по радиусам', 'радиус' in all_text.lower(), None)
    ck('PowerPoint: есть слайд конкурентной среды', 'Конкурентная среда' in all_text)
    ck('PowerPoint: есть слайд скоринга', 'скоринг' in all_text.lower())
    ck('PowerPoint: таблицы построены', tables >= 3, f'таблиц: {tables}')
    ck('PowerPoint: указан источник данных', 'OpenStreetMap' in all_text)
    ck('PowerPoint: формула скоринга приведена', '0,55' in all_text)
    ck('PowerPoint: есть слайд населения', 'Население вокруг точки' in all_text)
    ck('PowerPoint: есть разрез по районам города', 'районы города' in all_text.lower(), None)
    ck('PowerPoint: есть слайд источников и методики', 'Источники и методика' in all_text)
    ck('PowerPoint: указана плотность населения', 'плотность' in all_text.lower())

    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith('ppt/media/')]
        if media:
            png = z.read(media[0])
            ck('PowerPoint: снимок карты — настоящий PNG', png[:8] == b'\x89PNG\r\n\x1a\n', media[0])
            w = int.from_bytes(png[16:20], 'big')
            h = int.from_bytes(png[20:24], 'big')
            ck('PowerPoint: снимок карты нормального размера', w > 300 and h > 200, f'{w}×{h}')
        else:
            print('..  снимка карты нет — подложка не отдала тайлы с CORS, презентация собрана без картинки')


def main():
    d = sys.argv[1] if len(sys.argv) > 1 else 'exporttest'
    xs = glob.glob(os.path.join(d, '*.xlsx'))
    ps = glob.glob(os.path.join(d, '*.pptx'))
    ck('файл Excel найден', bool(xs), d)
    ck('файл PowerPoint найден', bool(ps), d)
    if xs:
        check_xlsx(xs[0])
    if ps:
        check_pptx(ps[0])
    print('\nПРОВАЛЕНО проверок: %d' % failed if failed else '\nВсе проверки пройдены')
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
